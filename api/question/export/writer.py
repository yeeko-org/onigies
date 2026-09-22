"""Block-level primitives over python-docx, bound to the template styles
created by `make_template.py`."""
from collections.abc import Iterable
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Emu, Pt, RGBColor
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

TEMPLATE = Path(__file__).resolve().parent / 'template.docx'

LOGO_PART = '/word/media/image2.png'
LOGO_WIDTH = Emu(996950)
LETTER_LIST = 'OnigiesLetter'
NOTE_ITEM = '- '

CHECKLIST_STYLE = 'Grid Table 1 Light Accent 4'
CHECKLIST_BANNER_FILL = 'FFDA66'
CHECKLIST_BANNER_BORDER = 'FFC000'
CHECKLIST_ACCENT = RGBColor(0x70, 0x30, 0xA0)
CHECKLIST_MARK_WIDTH = Emu(379095)
CHECKLIST_TEXT_WIDTH = Emu(6021705)
TEXT_WIDTH = Emu(6400800)


class DocxWriter:

    def __init__(self, template: Path = TEMPLATE) -> None:
        self.document = Document(str(template))
        self._logo = self._template_logo()
        self._letter_abstract_id = self._abstract_num_id(LETTER_LIST)

    def save(self) -> BytesIO:
        buffer = BytesIO()
        self.document.save(buffer)
        buffer.seek(0)
        return buffer

    def paragraph(
            self, text: str = '', style: str | None = None,
            page_break_before: bool = False) -> Paragraph:
        paragraph = self.document.add_paragraph(text, style=style)
        if page_break_before:
            paragraph.paragraph_format.page_break_before = True
        return paragraph

    def heading(
            self, text: str, level: int,
            page_break_before: bool = False) -> Paragraph:
        return self.paragraph(
            text, f'Heading {level}', page_break_before)

    def page_break(self) -> None:
        self.document.add_page_break()

    def logo(self, style: str = 'Part Logo') -> Paragraph:
        paragraph = self.paragraph(style=style)
        paragraph.add_run().add_picture(BytesIO(self._logo), width=LOGO_WIDTH)
        return paragraph

    def bullets(self, items: Iterable[str]) -> None:
        for item in items:
            self.paragraph(item, 'List Item')

    def lettered(self, items: Iterable[str]) -> None:
        """Lettered list restarting at «a.»: each call gets its own w:num,
        since a shared one would keep counting across observables."""
        numbering = self.document.part.numbering_part.element
        num = numbering.add_num(self._letter_abstract_id)
        num.add_lvlOverride(ilvl=0).add_startOverride(1)
        for item in items:
            paragraph = self.paragraph(item, 'Answer Option')
            num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
            num_pr.get_or_add_ilvl().val = 0
            num_pr.get_or_add_numId().val = num.numId

    def note(self, text: str) -> None:
        """Italic note; its «- » lines become list items, since a stored
        note is plain text edited in a textarea."""
        for line in text.splitlines():
            if not line.strip():
                continue
            if line.startswith(NOTE_ITEM):
                paragraph = self.paragraph(
                    line[len(NOTE_ITEM):], 'List Item')
            else:
                paragraph = self.paragraph(line)
            for run in paragraph.runs:
                run.italic = True

    def lines(self, items: Iterable[str]) -> None:
        for item in items:
            self.paragraph(item, 'Answer Line')

    def option_table(self, items: Iterable[str]) -> Table:
        table = self.document.add_table(rows=0, cols=1)
        table.style = self.document.styles['Option Table']
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for item in items:
            cell = table.add_row().cells[0]
            cell.width = TEXT_WIDTH
            cell.paragraphs[0].text = item
            cell.paragraphs[0].style = 'Table Text'
        return table

    def checklist_table(self, title: str, rows: Iterable[str]) -> Table:
        table = self.document.add_table(rows=1, cols=2)
        table.style = self.document.styles[CHECKLIST_STYLE]
        table.autofit = False
        table.columns[0].width = CHECKLIST_MARK_WIDTH
        table.columns[1].width = CHECKLIST_TEXT_WIDTH
        banner = table.rows[0].cells[0].merge(table.rows[0].cells[1])
        self._shade(banner, CHECKLIST_BANNER_FILL, CHECKLIST_BANNER_BORDER)
        self._cell_text(banner, title, size=16, center=True)
        for text in rows:
            mark_cell, text_cell = table.add_row().cells
            mark_cell.width = CHECKLIST_MARK_WIDTH
            text_cell.width = CHECKLIST_TEXT_WIDTH
            self._cell_text(text_cell, text, color=False)
        # The first column is the empty mark column: no style emphasis.
        table._tbl.tblPr.xpath('w:tblLook')[0].set(
            qn('w:firstColumn'), '0')
        return table

    def toc(self, placeholder: str, levels: str = '1-4') -> Paragraph:
        paragraph = self.paragraph()
        self._field_char(paragraph, 'begin')
        instr = paragraph.add_run()._r
        instr.append(parse_xml(
            f'<w:instrText {nsdecls("w")} xml:space="preserve">'
            f' TOC \\o "{levels}" \\h \\z \\u </w:instrText>'))
        self._field_char(paragraph, 'separate')
        paragraph.add_run(placeholder)
        self._field_char(paragraph, 'end')
        return paragraph

    def _cell_text(
            self, cell: _Cell, text: str, size: int | None = None,
            center: bool = False, color: bool = True) -> None:
        paragraph = cell.paragraphs[0]
        paragraph.style = 'Table Text'
        if center:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(text)
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = CHECKLIST_ACCENT

    @staticmethod
    def _shade(cell: _Cell, fill: str, top_border: str) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_pr.append(parse_xml(
            f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="12" '
            f'w:space="0" w:color="{top_border}"/></w:tcBorders>'))
        tc_pr.append(parse_xml(
            f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" '
            f'w:fill="{fill}"/>'))

    @staticmethod
    def _field_char(paragraph: Paragraph, kind: str) -> None:
        paragraph.add_run()._r.append(parse_xml(
            f'<w:fldChar {nsdecls("w")} w:fldCharType="{kind}"/>'))

    def _template_logo(self) -> bytes:
        for part in self.document.part.package.iter_parts():
            if part.partname == LOGO_PART:
                return part.blob
        raise LookupError(f'{LOGO_PART} missing from {TEMPLATE.name}')

    def _abstract_num_id(self, name: str) -> int:
        numbering = self.document.part.numbering_part.element
        for abstract in numbering.findall(qn('w:abstractNum')):
            name_el = abstract.find(qn('w:name'))
            if name_el is not None and name_el.get(qn('w:val')) == name:
                return int(abstract.get(qn('w:abstractNumId')))
        raise LookupError(f'numbering {name} missing from {TEMPLATE.name}')
